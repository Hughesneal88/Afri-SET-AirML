import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def format_run(run, text, bold=False, italic=False, font_size=11, color_rgb=(0,0,0)):
    run.text = text
    run.bold = bold
    run.italic = italic
    run.font.name = 'Calibri'
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color_rgb)

def parse_markdown_to_docx(md_filepath, docx_filepath):
    doc = Document()
    
    # Page Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
        
    with open(md_filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    def flush_table(lines_list):
        if not lines_list:
            return
        # Process markdown table
        rows_data = []
        for l in lines_list:
            l_str = l.strip()
            if not l_str.startswith('|'):
                continue
            cols = [c.strip() for c in l_str.split('|')[1:-1]]
            # Skip separator line (e.g. |---|---|)
            if all(set(c).issubset(set('-: ')) for c in cols):
                continue
            rows_data.append(cols)
            
        if not rows_data:
            return
            
        num_rows = len(rows_data)
        num_cols = max(len(r) for r in rows_data)
        
        tbl = doc.add_table(rows=num_rows, cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_idx, row in enumerate(rows_data):
            for c_idx, val in enumerate(row):
                if c_idx < num_cols:
                    cell = tbl.cell(r_idx, c_idx)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    
                    # Clean markdown formatting in cell text
                    clean_val = re.sub(r'[*_`$]', '', val)
                    r = p.add_run(clean_val)
                    if r_idx == 0:
                        format_run(r, clean_val, bold=True, font_size=10, color_rgb=(255, 255, 255))
                        set_cell_background(cell, "1F4E78")
                    else:
                        format_run(r, clean_val, bold=False, font_size=10, color_rgb=(50, 50, 50))
                        if r_idx % 2 == 1:
                            set_cell_background(cell, "F2F2F2")

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    for line in lines:
        line_str = line.strip()

        # Handle tables
        if line_str.startswith('|'):
            in_table = True
            table_lines.append(line_str)
            continue
        elif in_table:
            in_table = False
            flush_table(table_lines)
            table_lines = []

        if not line_str:
            continue

        # Horizontal Rule
        if line_str == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            continue

        # Headings
        if line_str.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(line_str[2:])
            format_run(r, line_str[2:], bold=True, font_size=20, color_rgb=(31, 78, 120))
        elif line_str.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(line_str[3:])
            format_run(r, line_str[3:], bold=True, font_size=15, color_rgb=(47, 85, 151))
        elif line_str.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(line_str[4:])
            format_run(r, line_str[4:], bold=True, font_size=13, color_rgb=(31, 78, 120))
        elif line_str.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(line_str[5:])
            format_run(r, line_str[5:], bold=True, font_size=11, color_rgb=(50, 50, 50))

        # Bullet List Items
        elif line_str.startswith('- ') or line_str.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            clean_text = line_str[2:]
            # Clean inline formatting markdown symbols
            clean_text = clean_text.replace('**', '').replace('__', '').replace('$', '')
            r = p.add_run(clean_text)
            format_run(r, clean_text, font_size=11, color_rgb=(40, 40, 40))

        # Numbered List Items
        elif re.match(r'^\d+\.\s', line_str):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            clean_text = re.sub(r'^\d+\.\s', '', line_str)
            clean_text = clean_text.replace('**', '').replace('__', '').replace('$', '')
            r = p.add_run(clean_text)
            format_run(r, clean_text, font_size=11, color_rgb=(40, 40, 40))

        # Blockquote / Notes
        elif line_str.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            clean_text = line_str[2:].replace('**', '').replace('$', '')
            r = p.add_run(clean_text)
            format_run(r, clean_text, italic=True, font_size=10.5, color_rgb=(80, 80, 80))

        # Standard Paragraphs
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            
            # Simple bold parsing
            parts = re.split(r'(\*\*.*?\*\*)', line_str)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    t = part[2:-2].replace('$', '')
                    r = p.add_run(t)
                    format_run(r, t, bold=True, font_size=11, color_rgb=(0, 0, 0))
                else:
                    t = part.replace('$', '')
                    r = p.add_run(t)
                    format_run(r, t, bold=False, font_size=11, color_rgb=(40, 40, 40))

    if in_table:
        flush_table(table_lines)

    doc.save(docx_filepath)
    print(f"Successfully converted: {os.path.basename(md_filepath)} -> {os.path.basename(docx_filepath)}")

if __name__ == '__main__':
    project_dir = os.path.dirname(os.path.abspath(__file__))
    output_folder = os.path.join(project_dir, "Thesis_Chapters_Word")
    os.makedirs(output_folder, exist_ok=True)

    chapters_map = [
        ("INTRODUCTION.md", "Chapter_1_Introduction.docx"),
        ("LITERATURE_REVIEW.md", "Chapter_2_Literature_Review.docx"),
        ("METHODOLOGY.md", "Chapter_3_Methodology.docx"),
        ("RESULTS_AND_DISCUSSION.md", "Chapter_4_Results_and_Discussion.docx"),
        ("CONCLUSION_AND_RECOMMENDATIONS.md", "Chapter_5_Conclusion_and_Recommendations.docx"),
    ]

    for md_file, docx_file in chapters_map:
        md_path = os.path.join(project_dir, md_file)
        docx_path = os.path.join(output_folder, docx_file)
        if os.path.exists(md_path):
            parse_markdown_to_docx(md_path, docx_path)

    # Also build single combined manuscript document
    combined_doc_path = os.path.join(output_folder, "Complete_Thesis_Manuscript.docx")
    master_doc = Document()
    for s in master_doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    title_p = master_doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(12)
    r_title = title_p.add_run("CALIBRATION OF LOW-COST PM2.5 SENSORS USING STACKED PHYSICS-INFORMED META-ENSEMBLES")
    format_run(r_title, r_title.text, bold=True, font_size=22, color_rgb=(31, 78, 120))
    
    subtitle_p = master_doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(24)
    r_sub = subtitle_p.add_run("Full Thesis Manuscript | Department of Physics, University of Ghana, Legon")
    format_run(r_sub, r_sub.text, italic=True, font_size=12, color_rgb=(100, 100, 100))
    master_doc.add_page_break()

    for idx, (md_file, _) in enumerate(chapters_map):
        md_path = os.path.join(project_dir, md_file)
        if os.path.exists(md_path):
            with open(md_path, 'r', encoding='utf-8') as f:
                content = f.read()
            # Append to master document
            temp_path = os.path.join(output_folder, f"temp_{idx}.docx")
            parse_markdown_to_docx(md_path, temp_path)
            temp_doc = Document(temp_path)
            
            for element in temp_doc.element.body:
                master_doc.element.body.append(element)
            
            os.remove(temp_path)
            if idx < len(chapters_map) - 1:
                master_doc.add_page_break()

    master_doc.save(combined_doc_path)
    print(f"Successfully generated Master Document: {os.path.basename(combined_doc_path)}")
